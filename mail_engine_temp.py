import openpyxl
from openpyxl.styles import Font, PatternFill
import smtplib
from email.message import EmailMessage
import time
import random
from datetime import datetime
import os
import database
import threading

class MailEngine:
    def __init__(self, excel_paths, smtp_server, port, email, password, limit, start_hour, end_hour, log_callback, progress_callback, finish_callback):
        self.excel_paths = excel_paths
        self.smtp_server = smtp_server
        self.port = port
        self.email = email
        self.password = password
        self.limit = limit
        self.start_hour = start_hour
        self.end_hour = end_hour
        self.log = log_callback
        self.update_progress = progress_callback
        self.on_finish = finish_callback
        
        self.is_running = False
        self.is_paused = False
        self.thread = None

    def start(self, col_email, col_name):
        if self.is_running:
            return
        self.is_running = True
        self.is_paused = False
        self.thread = threading.Thread(target=self._run, args=(col_email, col_name))
        self.thread.daemon = True
        self.thread.start()

    def pause(self):
        if self.is_running:
            self.is_paused = not self.is_paused
            if self.is_paused:
                self.log("[PAUZA] Kampanja je pauzirana.")
            else:
                self.log("[NASTAVAK] Kampanja se nastavlja...")

    def stop(self):
        if self.is_running:
            self.is_running = False
            self.log("[STOP] Zaustavljam kampanju, sačekaj...")

    def _connect(self):
        server = smtplib.SMTP_SSL(self.smtp_server, self.port)
        server.login(self.email, self.password)
        return server

    def _run(self, col_email, col_name):
        try:
            # Samo testiramo konekciju na pocetku
            server = self._connect()
            server.quit()
            self.log("✅ Kredencijali za Zoho su ispravni.")
        except Exception as e:
            self.log(f"❌ Greska pri konekciji: {e}")
            self.is_running = False
            self.on_finish()
            return

        sent_count = 0
        total_seconds = (self.end_hour - self.start_hour) * 3600
        average_sleep = total_seconds / max(1, self.limit)
        
        # Definisemo stil za "Poslato"
        sent_font = Font(color="006400", bold=True) # Tamno zelena, bold

        for excel_path in self.excel_paths:
            if not self.is_running:
                break
            if sent_count >= self.limit:
                break
                
            base_name = os.path.splitext(os.path.basename(excel_path))[0]
            templates_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Templates")
            template_path = None
            
            # Pokusavamo da nadjemo sablon cije ime se nalazi na pocetku imena baze (npr. Arhitektura_i_projektovanje)
            if os.path.exists(templates_dir):
                for t_file in os.listdir(templates_dir):
                    if t_file.endswith("_Template.txt"):
                        core_category = t_file.replace("_Template.txt", "")
                        if base_name.startswith(core_category):
                            template_path = os.path.join(templates_dir, t_file)
                            break
                            
            if not template_path:
                template_path = os.path.join(templates_dir, f"{base_name}_Template.txt")
            
            self.log(f"\n--- Otvaram fajl: {os.path.basename(excel_path)} ---")
            
            if not os.path.exists(template_path):
                self.log(f"GRESKA: Ne mogu da pronađem šablon za ovu bazu.")
                self.log(f"Očekivani fajl: {template_path}")
                self.log(f"Preskačem bazu {base_name}...")
                continue
                
            try:
                with open(template_path, 'r', encoding='utf-8') as tf:
                    template_content = tf.read().strip()
                tpl_lines = template_content.split('\n', 1)
                subject = tpl_lines[0].replace('Subject:', '').strip()
                body_template = tpl_lines[1].strip() if len(tpl_lines) > 1 else ""
            except Exception as e:
                self.log(f"GRESKA pri čitanju šablona: {e}")
                continue
                
            if not os.path.exists(excel_path):
                self.log(f"GRESKA: Fajl nije pronađen: {excel_path}")
                continue
                
            try:
                wb = openpyxl.load_workbook(excel_path)
            except Exception as e:
                self.log(f"GRESKA pri učitavanju Excela: {e}")
                continue

            header = {}
            for col_idx, cell in enumerate(ws[1], 1):
            for ws in wb.worksheets:
                if not self.is_running or self.sent_count >= self.limit:
                    break
                    if cell.value:
                        header[str(cell.value).strip()] = col_idx

                if col_email not in header:
                    self.log(f"GRESKA: Kolona '{col_email}' nije pronađena u prvom redu fajla.")
                    continue

                if col_name not in header:
                    self.log(f"Upozorenje: Kolona '{col_name}' nije pronađena, koristiće se 'kolege'.")

                # Dodajemo kolone ako ne postoje
                if 'Status' not in header:
                    new_col = ws.max_column + 1
                    ws.cell(row=1, column=new_col, value='Status')
                    header['Status'] = new_col

                if 'Datum Slanja' not in header:
                    new_col = ws.max_column + 1
                    ws.cell(row=1, column=new_col, value='Datum Slanja')
                    header['Datum Slanja'] = new_col
            
                unsent_indices = []
        
                for row_idx in range(2, ws.max_row + 1):
                    email_cell = ws.cell(row=row_idx, column=header[col_email]).value
                    status_cell = ws.cell(row=row_idx, column=header['Status']).value
            
                    email_val = str(email_cell).strip() if email_cell else ''
                    status = str(status_cell).strip() if status_cell else ''
            
                    if email_val.lower() in ['doslovmarko@gmail.com', 'bojan.mikulic@hotmail.com']:
                        unsent_indices.append(row_idx)
                        continue
            
                    if status != 'Poslato' and email_val and email_val.lower() != 'nan' and email_val.lower() != 'none':
                        if not database.is_email_sent(email_val):
                            unsent_indices.append(row_idx)
                        else:
                            ws.cell(row=row_idx, column=header['Status'], value='Poslato (Ranije)')
                    
                self.log(f"Pronadjeno {len(unsent_indices)} novih adresa u ovom fajlu.")
        
                if len(unsent_indices) == 0:
                    self.log("Sve adrese u ovom fajlu su već obrađene.")
                    try:
                        wb.save(excel_path)
                    except:
                        pass
                    continue

                to_send = unsent_indices
        
                for row_idx in to_send:
                    if sent_count >= self.limit:
                        self.log(f"\n[INFO] Dostignut dnevni limit od {self.limit} mejlova.")
                        break
                
                    while self.is_paused and self.is_running:
                        time.sleep(1)
                
                    if not self.is_running:
                        break
                
                    email_cell = ws.cell(row=row_idx, column=header[col_email]).value
                    email_addr = str(email_cell).strip() if email_cell else ''
            
                    company_name = "kolege"
                    if col_name in header:
                        company_cell = ws.cell(row=row_idx, column=header[col_name]).value
                        if company_cell and str(company_cell).strip().lower() not in ['nan', 'none', '']:
                            company_name = str(company_cell).strip()

                    hour = datetime.now().hour
                    is_working_hour = False
                    if self.start_hour <= self.end_hour:
                        is_working_hour = (self.start_hour <= hour < self.end_hour)
                    else:
                        is_working_hour = (hour >= self.start_hour or hour < self.end_hour)
                
                    if not is_working_hour:
                        self.log(f"Trenutno vreme ({hour}h) je van radnog vremena ({self.start_hour}-{self.end_hour}h).")
                        self.log("Zaustavljam kampanju za danas.")
                        self.is_running = False
                        break

                    self.log(f"[{sent_count+1}/{self.limit}] Saljem za: {company_name} ({email_addr})")
            
                    from email.utils import make_msgid
            
                    msg = EmailMessage()
                    # Personalizacija naslova i tela
                    msg['Subject'] = subject.replace("{company_name}", company_name)
                    msg['From'] = self.email
                    msg['To'] = email_addr
            
                    body_text = body_template.replace("{company_name}", company_name)
                    body_html = body_text.replace("\n", "<br>")
            
                    logo_path = r"E:\POSAO\3dmarch Elevate reality_Signature.png"
                    if os.path.exists(logo_path):
                        logo_cid = make_msgid()
                        # Add CID to HTML
                        body_html += f'<br><br><img src="cid:{logo_cid[1:-1]}" alt="3DMArch Logo">'
                
                        msg.set_content(body_text)
                        msg.add_alternative(body_html, subtype='html')
                
                        try:
                            with open(logo_path, 'rb') as img_file:
                                img_data = img_file.read()
                            msg.get_payload()[1].add_related(img_data, maintype='image', subtype='png', cid=logo_cid)
                        except Exception as img_e:
                            self.log(f"⚠️ Upozorenje: Nije moguce učitati sliku - {img_e}")
                    else:
                        msg.set_content(body_text)
                        msg.add_alternative(body_html, subtype='html')
            
                    try:
                        # Otvaramo konekciju tek sada, jer pauze traju po 15 minuta
                        server = self._connect()
                        server.send_message(msg)
                        server.quit()
                
                        self.log(f"✅ Uspesno poslato.")
                
                        if email_addr.lower() in ['doslovmarko@gmail.com', 'bojan.mikulic@hotmail.com']:
                            c_status = ws.cell(row=row_idx, column=header['Status'], value='Poslato (Test)')
                        else:
                            c_status = ws.cell(row=row_idx, column=header['Status'], value='Poslato')
                            database.add_sent_email(email_addr)
                    
                        # Primeni stil na celiju
                        c_status.font = sent_font
                    
                        ws.cell(row=row_idx, column=header['Datum Slanja'], value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
                        sent_count += 1
                    except Exception as e:
                        self.log(f"❌ Greska pri slanju: {e}")
                        ws.cell(row=row_idx, column=header['Status'], value='Greška')
                
                    try:
                        wb.save(excel_path)
                    except Exception as e:
                        self.log(f"⚠️ Upozorenje: Excel fajl je otvoren! Zatvori ga da bih mogao da sacuvam statuse.")

                    self.update_progress(sent_count / self.limit)

                    if not self.is_running or sent_count >= self.limit:
                        break
                
                    # Pause between emails
                    sleep_time = random.uniform(average_sleep * 0.8, average_sleep * 1.2)
                    mins, secs = divmod(sleep_time, 60)
                    self.log(f"⏳ Cekam {int(mins)} min {int(secs)} sek do sledeceg...")
            
                    slept = 0
                    while slept < sleep_time and self.is_running:
                        time.sleep(1)
                        if not self.is_paused:
                            slept += 1
        
                self.log(f"\n🎯 Kampanja završena! Ukupno poslato u ovoj sesiji: {sent_count}")
            self.is_running = False
        self.on_finish()
